# app/routes/docs_update.py
from fastapi import APIRouter, UploadFile, Form
from app.vectorstore.rag_store import RAGStore
import os
import pdfplumber
from docx import Document
import openpyxl

router = APIRouter()
rag_store = RAGStore()

def extract_pdf_text(file_path):
    text = ""
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + "\n"
    return text

def extract_docx_text(file_path):
    doc = Document(file_path)
    return "\n".join([p.text for p in doc.paragraphs])

def extract_excel_text(file_path):
    wb = openpyxl.load_workbook(file_path, data_only=True)
    text = ""
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            text += " ".join([str(cell) if cell else "" for cell in row]) + "\n"
    return text

@router.post("/docs/add")
async def add_doc(file: UploadFile = None, text: str = Form(None)):
    content_text = ""

    if file:
        contents = await file.read()
        ext = os.path.splitext(file.filename)[1].lower()
        tmp_path = f"temp_{file.filename}"
        with open(tmp_path, "wb") as f:
            f.write(contents)

        try:
            if ext == ".pdf":
                content_text = extract_pdf_text(tmp_path)
            elif ext == ".docx":
                content_text = extract_docx_text(tmp_path)
            elif ext in [".xls", ".xlsx"]:
                content_text = extract_excel_text(tmp_path)
            elif ext == ".txt":
                with open(tmp_path, "r", encoding="utf-8") as f:
                    content_text = f.read()
            else:
                return {"error": "Unsupported file type"}
        finally:
            os.remove(tmp_path)

    if text:
        content_text += "\n" + text

    if not content_text.strip():
        return {"error": "No content to add"}

    rag_store.add_document(content_text, filename=file.filename if file else "user_text")
    return {"status": "success", "message": "Document added"}
