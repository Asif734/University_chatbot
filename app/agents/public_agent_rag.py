import os
import pdfplumber
from docx import Document
import openpyxl
from app.vectorstore.vector_store import VectorStore

class PublicAgentRAG:
    def __init__(self):
        self.vector_store = VectorStore()
        # Optional: configure local LLM here
        # e.g., self.llm = LocalLLM(model_path="models/mistral-7b/")

    # ----------------- File Parsing -----------------
    def extract_text(self, file_path: str) -> str:
        ext = file_path.split(".")[-1].lower()
        text = ""
        if ext == "pdf":
            with pdfplumber.open(file_path) as pdf:
                text = "\n".join([page.extract_text() or "" for page in pdf.pages])
        elif ext == "docx":
            doc = Document(file_path)
            text = "\n".join([p.text for p in doc.paragraphs])
        elif ext in ["xls", "xlsx"]:
            wb = openpyxl.load_workbook(file_path)
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    text += " ".join([str(cell) for cell in row if cell]) + "\n"
        elif ext == "txt":
            with open(file_path, "r", encoding="utf-8") as f:
                text = f.read()
        return text

    # ----------------- Add Knowledge (from file) -----------------
    def add_document(self, file_path: str):
        text = self.extract_text(file_path)
        chunks = self.chunk_text(text)
        metadatas = [{"source": file_path}] * len(chunks)
        self.vector_store.add_texts(chunks, metadatas)

    # ----------------- Add Raw Text (manual input) -----------------
    def add_text(self, text: str, source: str = "manual_input"):
        chunks = self.chunk_text(text)
        metadatas = [{"source": source}] * len(chunks)
        self.vector_store.add_texts(chunks, metadatas)

    # ----------------- Chunking -----------------
    def chunk_text(self, text: str, chunk_size: int = 500, overlap: int = 50):
        words = text.split()
        chunks = []
        for i in range(0, len(words), chunk_size - overlap):
            chunks.append(" ".join(words[i:i+chunk_size]))
        return chunks

    # ----------------- Query -----------------
    def query(self, user_query: str, top_k: int = 5) -> str:
        results = self.vector_store.search(user_query, top_k)
        context = "\n".join([r[0] for r in results])

        prompt = f"""
        You are a multilingual assistant. Answer the user query based only on the context below.
        Do not publish what you are think, just share the exact answer.
        If the answer is not in the context, say you don't know.

        Context:
        {context}

        Question:
        {user_query}
        """

        try:
            import ollama
            response = ollama.chat(
                model="deepseek-r1:1.5b",  # replace with your local model
                messages=[{"role": "user", "content": prompt}]
            )
            return response["message"]["content"]
        except Exception as e:
            return f"[Fallback] Context only:\n{context}\n\nError calling LLM: {str(e)}"
